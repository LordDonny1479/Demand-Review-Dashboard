"""Reconcile a refresh to raw rows without using the dashboard aggregators."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import timedelta
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

import build_dashboard_data as mapping


ROOT = Path(__file__).resolve().parents[1]
MONTHS = mapping.MONTHS
ERRORS = []
COUNTS = Counter()
ROUNDING_BOUNDARIES = Counter()


def check(condition, message):
    COUNTS["checks"] += 1
    if not condition:
        ERRORS.append(message)


def round_case(value):
    return math.floor(value + 0.5) if value >= 0 else math.ceil(value - 0.5)


def check_rounded(actual, expected, message):
    target = round_case(expected)
    if actual != target and abs(abs(expected) % 1 - 0.5) < 0.00001:
        # The existing pipeline sums binary floats before whole-case rounding.
        valid = actual in {math.floor(expected), math.ceil(expected)}
        ROUNDING_BOUNDARIES["half_case_float_ties"] += int(valid)
        check(valid, message)
    else:
        check(actual == target, f"{message}: {actual} != {target}")


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


class EmbeddedData(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_data = False
        self.chunks = []

    def handle_starttag(self, tag, attrs):
        if tag == "script" and dict(attrs).get("id") == "dashboard-data":
            self.in_data = True

    def handle_endtag(self, tag):
        if tag == "script":
            self.in_data = False

    def handle_data(self, data):
        if self.in_data:
            self.chunks.append(data)


def verify_archive(path, expected):
    parsed = EmbeddedData()
    parsed.feed(path.read_text(encoding="utf-8"))
    check(json.loads("".join(parsed.chunks)) == expected, f"{path.name} embedded dashboard data exact match")


def date_counts(start, end):
    if end < start:
        start, end = end, start
    days = Counter()
    cursor = start
    while cursor <= end:
        days[cursor.month - 1] += 1
        cursor += timedelta(days=1)
    return days, sum(days.values())


def values():
    return {"base": [0.0] * 12, "comparison": [0.0] * 12}


def dimensions(row):
    return row["banner"], row["product_group"], row["mpg"]


def aggregates(details, period_labels, scopes):
    result = {scope: defaultdict(values) for scope in scopes}
    for row in details:
        period = "base" if row["period"] == period_labels["base"] else "comparison"
        month = MONTHS.index(row["month"])
        banner, group, mpg = dimensions(row)
        keys = [(None, None, None), (banner, None, None), (None, group, None),
                (None, group, mpg), (banner, group, None), (banner, group, mpg)]
        amount = row["exact_month_cases"]
        for scope, allowed in scopes.items():
            if allowed is not None and banner not in allowed:
                continue
            for key in keys:
                result[scope][key][period][month] += amount
    return result


def verify_row(row, expected, label):
    for period, month_field, total_field in [("base", "m25", "fy25"), ("comparison", "m26", "fy26")]:
        for index, amount in enumerate(expected[period]):
            check_rounded(row[month_field][index], amount, f"{label} {month_field} {MONTHS[index]}")
        check_rounded(row[total_field], sum(expected[period]), f"{label} {total_field}")
        check(row[month_field] == row[f"{period}_months"], f"{label} month alias")
        check(row[total_field] == row[f"{period}_total"], f"{label} total alias")
    COUNTS["table_rows"] += 1


def verify_table(rows, table, totals, label, fixed_banner=None):
    banner, group, mpg = fixed_banner, None, None
    for row in rows:
        if row.get("is_total"):
            key = (fixed_banner, None, None)
        elif table == "rollup_ret":
            if row.get("is_retailer"):
                banner, group, mpg = row["label"], None, None
            elif row.get("is_group"):
                group, mpg = row["label"], None
            elif row.get("is_mpg"):
                mpg = row["label"]
            key = (banner, group, mpg)
        else:
            if row.get("is_group"):
                banner, group, mpg = fixed_banner, row["label"], None
            elif row.get("is_mpg"):
                banner, mpg = fixed_banner, row["label"]
            elif row.get("is_retailer"):
                banner = row["label"]
            key = (banner, group, mpg)
        verify_row(row, totals.get(key, values()), f"{label}/{row['label']}")


def verify_promos(rows, totals, label):
    summed = defaultdict(lambda: {"m25": [0] * 12, "m26": [0] * 12, "fy25": 0, "fy26": 0})
    for row in rows:
        target = summed[dimensions(row)]
        for field in ("m25", "m26"):
            for month, amount in enumerate(row[field]):
                target[field][month] += amount
        for field in ("fy25", "fy26"):
            target[field] += row[field]
        check(bool(row["promo_id"]), f"{label} missing promo id")
        COUNTS["promo_rows"] += 1
    for key, actual in summed.items():
        expected = totals.get(key, values())
        for period, monthly, annual in [("base", "m25", "fy25"), ("comparison", "m26", "fy26")]:
            for month, amount in enumerate(expected[period]):
                check_rounded(actual[monthly][month], amount, f"{label} promo parent {key} {monthly}/{month}")
            check_rounded(actual[annual], sum(expected[period]), f"{label} promo parent {key}/{annual}")


def compare_baseline(previous, refreshed):
    checked = 0
    for mode in ("blended", "separate"):
        old = previous["RAW"]["comparisons"]["yoy"]["modes"][mode]
        new = refreshed["RAW"]["comparisons"]["mom"]["modes"][mode]
        for section in ("sales", "inventory"):
            a = old if section == "sales" else old["inventory"]
            b = new if section == "sales" else new["inventory"]
            for kind in ("retailer_totals", "promo_rows", "rollup_grp"):
                if kind not in a:
                    continue
                key_fn = (lambda row: row["label"]) if kind == "retailer_totals" else (lambda row: (*dimensions(row), row["promo_id"])) if kind == "promo_rows" else (lambda row: row.get("row_key", row["label"]))
                before = {key_fn(r): (*r["m26"], r["fy26"]) for r in a[kind] if kind != "rollup_grp" or not r.get("is_retailer")}
                after = {key_fn(r): (*r["m25"], r["fy25"]) for r in b[kind] if kind != "rollup_grp" or not r.get("is_retailer")}
                for key in before.keys() | after.keys():
                    check(before.get(key, (0,) * 13) == after.get(key, (0,) * 13), f"September baseline {mode}/{section}/{kind}/{key}")
                    checked += 1
    return checked


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--before", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "refresh-audit-2026-10-02.json")
    parser.add_argument("--before-html", type=Path)
    parser.add_argument("--after-html", type=Path)
    args = parser.parse_args()
    current = json.loads((ROOT / "public/data/promo-dashboard-data.json").read_text(encoding="utf-8"))
    previous = json.loads(args.before.read_text(encoding="utf-8"))
    if args.before_html:
        verify_archive(args.before_html, previous)
    if args.after_html:
        verify_archive(args.after_html, current)
    market_map, _ = mapping.load_market_map()
    products, item_by_id, pack_groups = mapping.load_products()
    raw_by_source = {}
    source_metadata = []
    for path in sorted({sheet["workbook"] for cfg in mapping.COMPARISON_CONFIGS.values() for sheet in cfg["sheets"]}):
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook["Rebates"]
        header = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
        check(header[2] == "Execution Start" and header[3] == "TLS Ship Start", f"{path.name} timing headers")
        check(header[7] == "Promo ID" and header[35] == "Fcst Inc Cases", f"{path.name} promo/volume headers")
        raw_by_source[path.name] = {index: row for index, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), 7) if any(v is not None for v in row)}
        source_metadata.append({"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "nonempty_rows": len(raw_by_source[path.name])})
        workbook.close()
    # Reuse the approved mapping rules only; arithmetic and aggregation below are independent.
    mapping_rows = []
    for cfg in mapping.COMPARISON_CONFIGS.values():
        mapping_rows.extend(mapping.raw_demand_rows(item_by_id, pack_groups, cfg["sheets"]))
    base_lookup = mapping.build_base_pack_lookup(products, mapping_rows)
    mapped = {(row["source_workbook"], row["source_row"]): row for row in mapping_rows}
    report = {"sources": source_metadata, "comparisons": {}, "spot_checks": []}
    for comparison, cfg in mapping.COMPARISON_CONFIGS.items():
        sales = read_csv(ROOT / "data" / f"promo-{comparison}-detail.csv")
        inventory = read_csv(ROOT / "data" / f"promo-{comparison}-inventory-detail.csv")
        excluded = read_csv(ROOT / "data" / f"promo-{comparison}-excluded-rows.csv")
        included = {(row["source_workbook"], int(row["source_row"])) for row in sales}
        excluded_keys = {(row["source_workbook"], int(row["source_row"])) for row in excluded}
        expected_included = set()
        source_set = {sheet["workbook"].name for sheet in cfg["sheets"]}
        all_keys = {(source, index) for source in source_set for index in raw_by_source[source]}
        for source, index in all_keys:
            raw = raw_by_source[source][index]
            rec = mapped[(source, index)]
            banner = market_map.get(mapping.clean(raw[0]))
            quantity = mapping.number(raw[35])
            allowed_statuses = {"Closed", "Committed"} if rec["year"] == 2025 else {"Closed", "Committed", "Planned"}
            valid = mapping.clean(raw[9]) in allowed_statuses and quantity > 0 and banner and banner != "Canada"
            product = rec["product_record"]
            valid = valid and product and (product.get("pack") or product.get("is_unspecified_pack")) and raw[2] and raw[5]
            if valid:
                expected_included.add((source, index))
        check(included == expected_included, f"{comparison} inclusion filters match raw")
        check(included.isdisjoint(excluded_keys), f"{comparison} included/excluded disjoint")
        check(included | excluded_keys == all_keys, f"{comparison} every raw row accounted for")
        check(len(excluded_keys) == len(excluded), f"{comparison} duplicate exclusions")
        expected_detail_keys = {"sales": Counter(), "inventory": Counter()}
        for source_key in expected_included:
            rec = mapped[source_key]
            conversion = mapping.conversion_for_row(rec, base_lookup)
            regular_mpg = mapping.pretty_pack_size(rec["product_record"]["pack_size"])
            definitions = [("blended", mapping.pretty_pack_size(c["base_pack_size"]), c["conversion_note"]) for c in conversion["components"]]
            definitions.append(("separate", regular_mpg, conversion["conversion_note"]))
            timing = [("sales", "", rec["execution_start"], rec["execution_end"])]
            if rec["tls_ship_start"]:
                timing.extend([("inventory", "Build", rec["tls_ship_start"], max(rec["tls_ship_start"], rec["execution_start"])), ("inventory", "Burn", rec["execution_start"], rec["execution_end"])])
            for table_name, movement, start, end in timing:
                months, _ = date_counts(start, end)
                for mode, mpg, note in definitions:
                    for month in months:
                        expected_detail_keys[table_name][(*source_key, mode, mpg, note, MONTHS[month], movement)] += 1
        for table_name, rows in [("sales", sales), ("inventory", inventory)]:
            actual_keys = Counter((r["source_workbook"], int(r["source_row"]), r["data_mode"], r["mpg"], r["conversion_note"], r["month"], r.get("movement_type", "")) for r in rows)
            check(actual_keys == expected_detail_keys[table_name], f"{comparison} {table_name} every component/month appears exactly once")
        period_by_source = {s["workbook"].name: s["label"] for s in cfg["sheets"]}
        inventory_balances = defaultdict(float)
        for table_name, rows in [("sales", sales), ("inventory", inventory)]:
            for detail in rows:
                source_key = (detail["source_workbook"], int(detail["source_row"]))
                rec = mapped[source_key]
                raw = raw_by_source[source_key[0]][source_key[1]]
                check(detail["period"] == period_by_source[source_key[0]], f"{comparison} {source_key} period")
                check(detail["banner"] == market_map.get(mapping.clean(raw[0])), f"{comparison} {source_key} banner")
                for field, col in [("product_id", 23), ("product", 24), ("promo_id", 7), ("promo_status", 9)]:
                    check(detail[field] == mapping.clean(raw[col]), f"{comparison} {source_key} {field}")
                quantity = mapping.number(raw[35])
                check(abs(float(detail["source_fcst_inc_cases"]) - quantity) < 0.000001, f"{source_key} raw incremental cases")
                conversion = mapping.conversion_for_row(rec, base_lookup)
                if detail["data_mode"] == "separate":
                    factor = 1.0
                    group = mapping.product_group_label(mapping.pretty_pack_size(rec["product_record"]["pack_size"]), rec["product_record"]["planner"], rec["product_record"]["segment"])
                    if conversion["unit_type"] == "DRP":
                        group = mapping.display_group_label(group)
                    check(detail["product_group"] == group, f"{source_key} separate group")
                    check(detail["mpg"] == mapping.pretty_pack_size(rec["product_record"]["pack_size"]), f"{source_key} separate MPG")
                else:
                    components = [c for c in conversion["components"] if mapping.pretty_pack_size(c["base_pack_size"]) == detail["mpg"] and c["conversion_note"] == detail["conversion_note"]]
                    check(len(components) == 1, f"{source_key} blended mapping component")
                    factor = components[0]["conversion"]
                    component = components[0]
                    group = component.get("product_group") or mapping.product_group_label(detail["mpg"], rec["product_record"]["planner"], rec["product_record"]["segment"])
                    check(detail["product_group"] == group, f"{source_key} blended group")
                mode_cases = quantity * factor
                check(abs(float(detail["mode_fcst_inc_cases"]) - mode_cases) < 0.000002, f"{source_key} converted volume")
                execution_start, execution_end = rec["execution_start"], rec["execution_end"]
                check(detail["execution_start"] == execution_start.isoformat() and detail["execution_end"] == execution_end.isoformat(), f"{source_key} execution dates")
                start, end, sign = execution_start, execution_end, 1
                day_fields = ("execution_days_in_month", "execution_days_total")
                if table_name == "inventory":
                    sign = 1 if detail["movement_type"] == "Build" else -1
                    if sign == 1:
                        start = rec["tls_ship_start"]
                        end = max(start, execution_start)
                    day_fields = ("movement_days_in_month", "movement_days_total")
                    check(int(detail["movement_sign"]) == sign, f"{source_key} movement sign")
                    check(detail["tls_ship_start"] == rec["tls_ship_start"].isoformat(), f"{source_key} ship start")
                    check(detail["timing_start"] == start.isoformat() and detail["timing_end"] == end.isoformat(), f"{source_key} movement timing")
                month_days, total_days = date_counts(start, end)
                month = MONTHS.index(detail["month"])
                check(int(detail[day_fields[0]]) == month_days[month] and int(detail[day_fields[1]]) == total_days, f"{source_key} {table_name} inclusive day counts")
                weight = month_days[month] / total_days
                check(abs(float(detail["prorate_weight"]) - weight) < 0.00000001, f"{source_key} proration")
                expected = mode_cases * weight * sign
                check(abs(float(detail["month_cases"]) - expected) < 0.000002, f"{source_key} {table_name} month cases")
                detail["exact_month_cases"] = expected
                COUNTS[f"{table_name}_detail_rows"] += 1
                if table_name == "inventory":
                    inventory_balances[(*source_key, detail["data_mode"], detail["product_group"], detail["mpg"], detail["conversion_note"])] += expected
        for key, amount in inventory_balances.items():
            check(abs(amount) < 0.00001, f"{comparison} inventory lifecycle {key}")
        for mode, dashboard in current["RAW"]["comparisons"][comparison]["modes"].items():
            scopes = {"all": None, "regular": set(current["RAW"]["all_banner_order"]) - {"Amazon", "Costco", "Canada"}, "non_mulo": {"Amazon", "Costco"}}
            for table_name, details in [("sales", sales), ("inventory", inventory)]:
                selected = [row for row in details if row["data_mode"] == mode]
                grouped = aggregates(selected, cfg["period_labels"], scopes)
                data = dashboard if table_name == "sales" else dashboard["inventory"]
                for table in ("rollup_ret", "rollup_grp", "rollup_segment"):
                    verify_table(data[table], table, grouped["regular"], f"{comparison}/{mode}/{table_name}/{table}")
                for row in data.get("retailer_totals", []):
                    verify_row(row, grouped["all"][(row["label"], None, None)], f"{comparison}/{mode}/{table_name}/tile/{row['label']}")
                for banner, rows in data.get("retailers", {}).items():
                    verify_table(rows, "product", grouped["all"], f"{comparison}/{mode}/{table_name}/{banner}", fixed_banner=banner)
                verify_promos(data["promo_rows"], grouped["all"], f"{comparison}/{mode}/{table_name}")
                for period, field in [("base", "fy25"), ("comparison", "fy26")]:
                    check_rounded(data["stats"][field], sum(grouped["regular"][(None, None, None)][period]), f"{comparison}/{mode}/{table_name}/stats/{field}")
                check(data["stats"]["delta"] == data["stats"]["fy26"] - data["stats"]["fy25"], f"{comparison}/{mode}/{table_name}/delta")
                if table_name == "sales":
                    for table in ("rollup_ret", "rollup_grp", "rollup_segment"):
                        verify_table(data["non_mulo"][table], table, grouped["non_mulo"], f"{comparison}/{mode}/non_mulo/{table}")
        report["comparisons"][comparison] = {"raw_rows": len(all_keys), "included_source_rows": len(included), "excluded_source_rows": len(excluded), "excluded_reasons": dict(Counter(r["reason"] for r in excluded)), "blended_totals": current["RAW"]["comparisons"][comparison]["modes"]["blended"]["stats"], "separate_totals": current["RAW"]["comparisons"][comparison]["modes"]["separate"]["stats"]}
        if comparison == "yoy":
            report["unmapped_retailer_exceptions"] = []
            for excluded_row in excluded:
                if excluded_row["reason"] != "Market is outside the supplied market list":
                    continue
                key = (excluded_row["source_workbook"], int(excluded_row["source_row"]))
                rec = mapped[key]
                conversion = mapping.conversion_for_row(rec, base_lookup)
                report["unmapped_retailer_exceptions"].append({"market": rec["market"], "promo_id": rec["promo_id"], "product_id": rec["product_id"], "source_volume": rec["forecast_incremental_cases"], "unit_type": conversion["unit_type"], "blended_cases_if_mapped": rec["forecast_incremental_cases"] * conversion["conversion"], "source_workbook": key[0], "source_row": key[1], "status": "Excluded under the existing Market List; mapping requires user confirmation"})
        for retailer in ("Walmart", "Canadian Tire", "Amazon", "Costco"):
            row = next(r for r in current["RAW"]["comparisons"][comparison]["modes"]["blended"]["retailer_totals"] if r["label"] == retailer)
            report["spot_checks"].append({"comparison": comparison, "retailer": retailer, "base_cases": row["fy25"], "current_cases": row["fy26"], "change": row["fy26"] - row["fy25"], "current_months": row["m26"]})
    report["baseline_vectors_checked"] = compare_baseline(previous, current)
    check(json.loads((ROOT / "data/promo-yoy-dashboard.json").read_text(encoding="utf-8")) == current["RAW"], "generated YoY JSON matches public payload")
    check(json.loads((ROOT / "data/promo-mom-dashboard.json").read_text(encoding="utf-8")) == current["RAW"]["comparisons"]["mom"], "generated MoM JSON matches public payload")
    check(json.loads((ROOT / "data/dashboard-summary.json").read_text(encoding="utf-8")) == current["META"], "summary matches public payload")
    for comparison in ("yoy", "mom"):
        check(current["META"]["comparisons"][comparison]["unconverted_display_products"] == 0, f"{comparison} all displays converted")
    prior_conversion_file = args.before.parent / "september-conversions-before.csv"
    if prior_conversion_file.exists():
        prior_conversions = read_csv(prior_conversion_file)
        new_conversions = read_csv(ROOT / "data/display-conversion-audit.csv")
        key = lambda r: (r["product_id"], r["product"], r["blended_mpg"], r["separate_mpg"])
        before = {key(r): (r["cases_per_display"], r["blended_product_group"]) for r in prior_conversions}
        after = {key(r): (r["cases_per_display"], r["blended_product_group"]) for r in new_conversions}
        for common in before.keys() & after.keys():
            check(before[common] == after[common], f"approved conversion unchanged {common}")
        report["existing_display_components_unchanged"] = len(before.keys() & after.keys())
    for original, copied in [("EXCEXP_TLS_000JJKSJR_ExcelExport.xlsx", mapping.OCTOBER_2026_XLSX), ("EXCEXP_TLS_000JJKSK3_ExcelExport.xlsx", mapping.YOY_2025_XLSX)]:
        check(hashlib.sha256((Path.home() / "Downloads" / original).read_bytes()).digest() == hashlib.sha256(copied.read_bytes()).digest(), f"{original} raw copy unchanged")
    report.update({"checks": dict(COUNTS), "rounding_boundaries": dict(ROUNDING_BOUNDARIES), "errors": ERRORS, "result": "PASS" if not ERRORS else "FAIL"})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps({"result": report["result"], "checks": dict(COUNTS), "baseline_vectors_checked": report["baseline_vectors_checked"], "errors": ERRORS[:15], "audit": str(args.output)}, indent=2))
    if ERRORS:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
